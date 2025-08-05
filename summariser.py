import os
import docx
import json
import time
from concurrent.futures import ThreadPoolExecutor
from PyPDF2 import PdfReader
from transformers import PegasusTokenizer, AutoModelForSeq2SeqLM
from fpdf import FPDF
from keybert import KeyBERT

# Load models
tokenizer = PegasusTokenizer.from_pretrained("nsi319/legal-pegasus")
model = AutoModelForSeq2SeqLM.from_pretrained("nsi319/legal-pegasus")
keybert_model = KeyBERT()

# Text extraction functions remain the same
def extract_text_from_txt(file_path):
    with open(file_path, 'r', encoding='utf-8') as file:
        return file.read()

def extract_text_from_pdf(file_path):
    pdf_reader = PdfReader(file_path)
    text = ""
    for page in pdf_reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def extract_text_from_docx(file_path):
    doc = docx.Document(file_path)
    return "\n".join(para.text for para in doc.paragraphs)

# Split text into sections based on legal terms
def split_into_sections(text):
    sections = {}
    current_section = "Introduction"
    sections[current_section] = []

    for line in text.splitlines():
        if any(heading in line for heading in [
            "DEFINITIONS", "PAYMENT OF FEES", "LICENSE", 
            "CONFIDENTIALITY", "LIMITATION OF LIABILITY", 
            "INDEMNITIES", "TERMINATION", "WARRANTIES", "GOVERNING LAW"
        ]):
            current_section = line.strip()
            sections[current_section] = []
        else:
            sections[current_section].append(line)
    
    for section, lines in sections.items():
        sections[section] = " ".join(lines).strip()
    
    return sections

# Extract keywords using KeyBERT
def extract_keywords(text, num_keywords=10):
    keywords = keybert_model.extract_keywords(text, top_n=num_keywords)
    return ", ".join([kw[0] for kw in keywords])

# Summarize a single text chunk
def summarize_chunk(chunk, min_length=300, max_length=800):
    key_phrases = extract_keywords(chunk)
    chunk_with_keywords = key_phrases + " " + chunk

    inputs = tokenizer.encode(chunk_with_keywords, return_tensors="pt", truncation=True)
    outputs = model.generate(
        inputs,
        max_length=max_length,
        min_length=min_length,
        num_beams=4,
        length_penalty=2.0,
        no_repeat_ngram_size=3
    )
    return tokenizer.decode(outputs[0], skip_special_tokens=True)

# Summarize all sections with customizable length
def summarize_sections(sections, min_length=150, max_length=300, chunk_size=2000):
    summaries = {}
    for section, content in sections.items():
        if len(content.split()) > chunk_size:
            chunks = [content[i:i + chunk_size] for i in range(0, len(content.split()), chunk_size)]
        else:
            chunks = [content]

        section_summaries = []
        with ThreadPoolExecutor() as executor:
            section_summaries = list(executor.map(lambda chunk: summarize_chunk(chunk, min_length, max_length), chunks))
        combined_section_summary = ' '.join(section_summaries)
        if len(combined_section_summary.split()) > max_length * 1.5:
            inputs = tokenizer.encode(combined_section_summary, return_tensors="pt", truncation=True)
            outputs = model.generate(
                inputs,
                max_length=int(max_length * 0.75),
                min_length=min_length,
                num_beams=4,
                length_penalty=2.5,
                no_repeat_ngram_size=3
            )
            combined_section_summary = tokenizer.decode(outputs[0], skip_special_tokens=True)

        # Enhanced formatting with better structure
        formatted_summary = format_section_summary(combined_section_summary, content)
        summaries[section] = formatted_summary
    
    return summaries

def format_section_summary(summary_text, original_content):
    """Enhanced formatting for better readability."""
    
    # Split into sentences and clean them
    sentences = [s.strip() for s in summary_text.split('.') if s.strip()]
    
    # Group sentences into logical points
    formatted_points = []
    current_point = ""
    
    for sentence in sentences:
        if len(current_point) + len(sentence) < 150:  # Combine short sentences
            current_point += sentence + ". "
        else:
            if current_point:
                formatted_points.append(current_point.strip())
            current_point = sentence + ". "
    
    if current_point:
        formatted_points.append(current_point.strip())
    
    # Format as numbered list for better structure
    formatted_text = ""
    for i, point in enumerate(formatted_points, 1):
        formatted_text += f"{i}. {point}\n"
    
    # Add key terms section
    keywords = extract_keywords(original_content, num_keywords=5)
    if keywords:
        formatted_text += f"\n📋 Key Terms: {keywords}\n"
    
    return formatted_text.strip()

def compile_final_summary(summaries):
    """Compile summaries with enhanced formatting and structure."""
    
    # Create a professional header
    header = "=" * 60 + "\n"
    header += "📄 TERMS & CONDITIONS SUMMARY\n"
    header += "=" * 60 + "\n\n"
    
    # Add executive summary if multiple sections
    if len(summaries) > 1:
        header += "🔍 EXECUTIVE OVERVIEW\n"
        header += "-" * 25 + "\n"
        header += "This document has been analyzed and summarized into key sections below.\n"
        header += "Each section contains the most important points in plain language.\n\n"
    
    # Format each section with better structure
    formatted_sections = []
    
    for section_name, summary in summaries.items():
        section_header = f"📋 {section_name.upper()}\n"
        section_header += "-" * (len(section_name) + 4) + "\n"
        
        formatted_section = section_header + summary + "\n"
        formatted_sections.append(formatted_section)
    
    # Add footer with disclaimer
    footer = "\n" + "=" * 60 + "\n"
    footer += "⚠️  IMPORTANT DISCLAIMER\n"
    footer += "=" * 60 + "\n"
    footer += "This is an AI-generated summary for informational purposes only.\n"
    footer += "Please refer to the original document for complete legal terms.\n"
    footer += "Consult with legal professionals for important decisions.\n"
    
    final_summary = header + "\n\n".join(formatted_sections) + footer
    return final_summary

# Helper function to format text for PDF
def clean_text_for_pdf(text):
    return text.encode('latin-1', 'replace').decode('latin-1')

# Enhanced PDF formatting with better styling
def save_summary_as_pdf(summary, output_path="summary_output_legal.pdf"):
    """Create a professionally formatted PDF with enhanced styling."""
    
    summary = clean_text_for_pdf(summary)
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    
    # Add header with styling
    pdf.set_font("Arial", style="B", size=16)
    pdf.set_text_color(0, 51, 102)  # Dark blue
    pdf.cell(0, 10, "Terms & Conditions Summary", ln=True, align="C")
    pdf.ln(5)
    
    # Add generation date
    pdf.set_font("Arial", size=10)
    pdf.set_text_color(128, 128, 128)  # Gray
    from datetime import datetime
    pdf.cell(0, 5, f"Generated on: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}", ln=True, align="C")
    pdf.ln(10)
    
    # Reset text color
    pdf.set_text_color(0, 0, 0)
    
    # Process content with enhanced formatting
    lines = summary.split('\n')
    
    for line in lines:
        line = line.strip()
        if not line:
            pdf.ln(3)
            continue
            
        # Handle different types of content
        if line.startswith('='):
            # Main headers
            pdf.set_font("Arial", style="B", size=14)
            pdf.set_text_color(0, 51, 102)
            if not line.replace('=', '').strip():
                pdf.ln(2)
            else:
                header_text = line.replace('=', '').strip()
                if header_text:
                    pdf.cell(0, 8, header_text, ln=True, align="C")
                    pdf.ln(3)
        
        elif line.startswith('📋') or line.startswith('🔍'):
            # Section headers
            pdf.set_font("Arial", style="B", size=12)
            pdf.set_text_color(0, 102, 51)  # Dark green
            clean_line = line.encode('ascii', 'ignore').decode('ascii')
            pdf.cell(0, 8, clean_line, ln=True)
            pdf.ln(2)
        
        elif line.startswith('-'):
            # Dashed separators
            pdf.set_font("Arial", size=10)
            pdf.set_text_color(128, 128, 128)
            pdf.cell(0, 4, line, ln=True)
            pdf.ln(1)
        
        elif line.startswith('⚠️'):
            # Warning/disclaimer
            pdf.set_font("Arial", style="B", size=10)
            pdf.set_text_color(204, 102, 0)  # Orange
            clean_line = line.encode('ascii', 'ignore').decode('ascii')
            pdf.multi_cell(0, 5, clean_line)
            pdf.ln(2)
        
        elif any(line.startswith(str(i) + '.') for i in range(1, 10)):
            # Numbered points
            pdf.set_font("Arial", size=11)
            pdf.set_text_color(0, 0, 0)
            pdf.multi_cell(0, 6, line, align="J")
            pdf.ln(2)
        
        elif line.startswith('Key Terms:') or line.startswith('📋 Key Terms:'):
            # Key terms
            pdf.set_font("Arial", style="I", size=10)
            pdf.set_text_color(51, 51, 153)  # Blue
            clean_line = line.encode('ascii', 'ignore').decode('ascii')
            pdf.multi_cell(0, 5, clean_line)
            pdf.ln(3)
        
        else:
            # Regular text
            pdf.set_font("Arial", size=11)
            pdf.set_text_color(0, 0, 0)
            pdf.multi_cell(0, 6, line, align="J")
            pdf.ln(2)
    
    # Add footer
    pdf.ln(10)
    pdf.set_font("Arial", style="I", size=9)
    pdf.set_text_color(128, 128, 128)
    pdf.multi_cell(0, 4, "This summary was generated using AI technology. Please consult the original document and legal professionals for complete information.")
    
    pdf.output(output_path)

# Function for customizable summary length
def summarize_with_custom_length(text, custom_min_length, custom_max_length):
    sections = split_into_sections(text)
    summaries = summarize_sections(sections, min_length=custom_min_length, max_length=custom_max_length)
    final_summary = compile_final_summary(summaries)
    return final_summary

# Feedback collection function
def store_feedback(feedback_text, feedback_file="feedback.json"):
    with open(feedback_file, "a") as file:
        json.dump({"feedback": feedback_text}, file)
        file.write("\n")
